"""Real tests: hand-verified DCF arithmetic + xlsx integrity."""
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from main import app, build_dcf, DCFInput

client = TestClient(app)

SIMPLE = {
    "initial_revenue": 100.0,
    "growth_rates": [0.0],
    "ebit_margin": 0.20,
    "tax_rate": 0.25,
    "capex_pct_revenue": 0.0,
    "nwc_pct_revenue": 0.0,
    "discount_rate": 0.10,
    "terminal_growth_rate": 0.02,
    "net_debt": 0.0,
    "shares_outstanding": 10.0,
}


def test_hand_verified_single_year():
    # Hand calc: NOPAT = 100*0.2*0.75 = 15; FCF = 15
    # PV(FCF) = 15/1.1 = 13.636363...
    # TV = 15*1.02/(0.10-0.02) = 191.25; PV(TV) = 191.25/1.1 = 173.863636...
    # EV = 187.5 ; per share = 18.75
    m = build_dcf(DCFInput(**SIMPLE))
    assert m["years"][0]["fcf"] == 15.0
    assert abs(m["pv_fcf_sum"] - 15 / 1.1) < 1e-9
    assert abs(m["terminal_value"] - 191.25) < 1e-9
    assert abs(m["pv_terminal_value"] - 191.25 / 1.1) < 1e-9
    assert abs(m["enterprise_value"] - 187.5) < 1e-9
    assert abs(m["value_per_share"] - 18.75) < 1e-9


def test_ev_identity_multi_year():
    inp = DCFInput(
        initial_revenue=500.0, growth_rates=[0.10, 0.08, 0.06, 0.04, 0.03],
        ebit_margin=0.25, tax_rate=0.25, capex_pct_revenue=0.06,
        nwc_pct_revenue=0.12, discount_rate=0.09, terminal_growth_rate=0.025,
        net_debt=120.0, shares_outstanding=50.0,
    )
    m = build_dcf(inp)
    # EV must equal PV(FCFs) + PV(TV) exactly
    assert abs(m["enterprise_value"] - (m["pv_fcf_sum"] + m["pv_terminal_value"])) < 1e-9
    assert abs(m["equity_value"] - (m["enterprise_value"] - 120.0)) < 1e-9
    assert abs(m["value_per_share"] - m["equity_value"] / 50.0) < 1e-9
    # FCF identity per year
    for y in m["years"]:
        assert abs(y["fcf"] - (y["nopat"] - y["capex"] - y["delta_nwc"])) < 1e-9
    # Revenue compounds correctly
    assert abs(m["years"][0]["revenue"] - 550.0) < 1e-9
    assert abs(m["years"][1]["revenue"] - 550.0 * 1.08) < 1e-9


def test_terminal_growth_validation():
    bad = dict(SIMPLE, terminal_growth_rate=0.10, discount_rate=0.09)
    r = client.post("/dcf", json=bad)
    assert r.status_code == 400


def test_xlsx_download_opens():
    r = client.post("/dcf/xlsx", json=SIMPLE)
    assert r.status_code == 200
    assert "spreadsheetml.sheet" in r.headers["content-type"]
    wb = load_workbook(filename=__import__("io").BytesIO(r.content))
    ws = wb["DCF Model"]
    assert ws["A1"].value == "DCF Valuation Model"
    # find the valuation summary value-per-share cell and check the number
    found = None
    for row in ws.iter_rows(values_only=False):
        for cell in row:
            if cell.value == "Value per share":
                found = ws.cell(row=cell.row, column=2).value
    assert found is not None and abs(found - 18.75) < 1e-6
    # yearly FCF row present
    assert ws.cell(row=17, column=8).value == 15.0 or any(
        ws.cell(row=rr, column=8).value == 15.0 for rr in range(1, 40)
    )


def test_json_endpoint():
    r = client.post("/dcf", json=SIMPLE)
    assert r.status_code == 200
    assert abs(r.json()["value_per_share"] - 18.75) < 1e-9
