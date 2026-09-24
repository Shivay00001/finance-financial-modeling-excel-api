"""Real DCF (discounted cash flow) financial model with Excel export.

Model (all real arithmetic, no stubs):
  revenue_t      = revenue_{t-1} * (1 + g_t)
  ebit_t         = revenue_t * ebit_margin
  nopat_t        = ebit_t * (1 - tax_rate)
  capex_t        = revenue_t * capex_pct_revenue
  nwc_t          = revenue_t * nwc_pct_revenue        (working-capital level)
  fcf_t          = nopat_t - capex_t - (nwc_t - nwc_{t-1})
  df_t           = 1 / (1 + r)^t
  pv_fcf         = sum(fcf_t * df_t)
  TV             = fcf_n * (1 + g_terminal) / (r - g_terminal)   (Gordon growth)
  enterprise_value = pv_fcf + TV * df_n
  equity_value   = enterprise_value - net_debt
  value_per_share = equity_value / shares_outstanding
"""

from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

app = FastAPI(title="Financial Modeling Excel API — DCF")


class DCFInput(BaseModel):
    initial_revenue: float = Field(..., gt=0, description="Year-0 revenue")
    growth_rates: list[float] = Field(..., min_length=1, description="Per-year revenue growth rates, e.g. [0.10, 0.08, 0.05]")
    ebit_margin: float = Field(..., description="EBIT as fraction of revenue, e.g. 0.20")
    tax_rate: float = Field(..., ge=0, lt=1)
    capex_pct_revenue: float = Field(0.05, ge=0)
    nwc_pct_revenue: float = Field(0.10, ge=0)
    discount_rate: float = Field(..., gt=0, description="WACC, e.g. 0.10")
    terminal_growth_rate: float = Field(..., ge=0, description="Perpetual growth, must be < discount_rate")
    net_debt: float = Field(0.0)
    shares_outstanding: float = Field(1.0, gt=0)


def build_dcf(inp: DCFInput) -> dict:
    if inp.terminal_growth_rate >= inp.discount_rate:
        raise HTTPException(400, "terminal_growth_rate must be < discount_rate")
    n = len(inp.growth_rates)
    r = inp.discount_rate
    years = []
    revenue_prev = inp.initial_revenue
    nwc_prev = inp.initial_revenue * inp.nwc_pct_revenue
    pv_fcf_total = 0.0
    for t, g in enumerate(inp.growth_rates, start=1):
        revenue = revenue_prev * (1 + g)
        ebit = revenue * inp.ebit_margin
        nopat = ebit * (1 - inp.tax_rate)
        capex = revenue * inp.capex_pct_revenue
        nwc = revenue * inp.nwc_pct_revenue
        delta_nwc = nwc - nwc_prev
        fcf = nopat - capex - delta_nwc
        df = 1.0 / (1 + r) ** t
        pv = fcf * df
        pv_fcf_total += pv
        years.append({
            "year": t, "growth": g, "revenue": revenue, "ebit": ebit,
            "nopat": nopat, "capex": capex, "delta_nwc": delta_nwc,
            "fcf": fcf, "discount_factor": df, "pv_fcf": pv,
        })
        revenue_prev, nwc_prev = revenue, nwc
    fcf_n = years[-1]["fcf"]
    tv = fcf_n * (1 + inp.terminal_growth_rate) / (r - inp.terminal_growth_rate)
    df_n = 1.0 / (1 + r) ** n
    pv_tv = tv * df_n
    ev = pv_fcf_total + pv_tv
    equity = ev - inp.net_debt
    return {
        "years": years,
        "terminal_value": tv,
        "pv_fcf_sum": pv_fcf_total,
        "pv_terminal_value": pv_tv,
        "enterprise_value": ev,
        "equity_value": equity,
        "value_per_share": equity / inp.shares_outstanding,
        "inputs": inp.model_dump(),
    }


def build_workbook(model: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "DCF Model"
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")
    money = '#,##0.00'

    ws["A1"] = "DCF Valuation Model"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "INPUTS"
    ws["A3"].font = hdr
    ws["A3"].fill = fill
    inputs = model["inputs"]
    labels = [("Initial revenue", inputs["initial_revenue"]),
              ("EBIT margin", inputs["ebit_margin"]),
              ("Tax rate", inputs["tax_rate"]),
              ("Capex % revenue", inputs["capex_pct_revenue"]),
              ("NWC % revenue", inputs["nwc_pct_revenue"]),
              ("Discount rate (WACC)", inputs["discount_rate"]),
              ("Terminal growth", inputs["terminal_growth_rate"]),
              ("Net debt", inputs["net_debt"]),
              ("Shares outstanding", inputs["shares_outstanding"])]
    for i, (label, val) in enumerate(labels, start=4):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = val
        ws[f"B{i}"].number_format = money

    start = 4 + len(labels) + 2
    ws[f"A{start}"] = "YEARLY FREE CASH FLOW"
    ws[f"A{start}"].font = hdr
    ws[f"A{start}"].fill = fill
    cols = ["Year", "Growth", "Revenue", "EBIT", "NOPAT", "Capex",
            "dNWC", "FCF", "Disc. factor", "PV of FCF"]
    keys = ["year", "growth", "revenue", "ebit", "nopat", "capex",
            "delta_nwc", "fcf", "discount_factor", "pv_fcf"]
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=start + 1, column=j, value=c)
        cell.font = hdr
        cell.fill = fill
    for i, y in enumerate(model["years"]):
        for j, k in enumerate(keys, start=1):
            cell = ws.cell(row=start + 2 + i, column=j, value=y[k])
            if j >= 3:
                cell.number_format = money

    srow = start + 2 + len(model["years"]) + 1
    ws[f"A{srow}"] = "VALUATION SUMMARY"
    ws[f"A{srow}"].font = hdr
    ws[f"A{srow}"].fill = fill
    summary = [("PV of forecast FCF", model["pv_fcf_sum"]),
               ("Terminal value", model["terminal_value"]),
               ("PV of terminal value", model["pv_terminal_value"]),
               ("Enterprise value", model["enterprise_value"]),
               ("Equity value", model["equity_value"]),
               ("Value per share", model["value_per_share"])]
    for i, (label, val) in enumerate(summary, start=srow + 1):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = val
        ws[f"B{i}"].number_format = money
        if "Enterprise" in label or "per share" in label:
            ws[f"A{i}"].font = Font(bold=True)
            ws[f"B{i}"].font = Font(bold=True)
    for col in range(1, 11):
        ws.column_dimensions[get_column_letter(col)].width = 16
    return wb


@app.get("/health")
def health():
    return {"status": "ok", "model": "dcf"}


@app.post("/dcf")
def dcf_json(inp: DCFInput):
    return build_dcf(inp)


@app.post("/dcf/xlsx")
def dcf_xlsx(inp: DCFInput):
    model = build_dcf(inp)
    wb = build_workbook(model)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=dcf_model.xlsx"},
    )
